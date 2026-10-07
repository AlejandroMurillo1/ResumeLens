
import docx
import pytest

from src.core.ingestion.docx_loader import DocxLoader
from src.core.ingestion.loader_factory import LoaderFactory
from src.core.ingestion.txt_loader import TxtLoader
from src.core.ingestion.md_loader import MdLoader
from src.core.ingestion.pdf_loader import PdfLoader
from src.core.ingestion.file_loader import IngestionError, UnsupportedFileFormatError

class  TestTxtLoader:
    def test_valid_file_with_content_returns_content(self, tmp_path):
        path = tmp_path / "cv.txt"
        path.write_text("John Doe\nSkills: Python, SQL", encoding="utf-8")

        result = TxtLoader().load(path)

        assert result == "John Doe\nSkills: Python, SQL"

    def test_empty_file_returns_empty_string(self, tmp_path):
        path = tmp_path / "empty.txt"
        path.write_text("", encoding="utf-8")

        assert TxtLoader().load(path) == ""

    def test_non_utf8_encoding_raises_ingestion_error(self, tmp_path):
        path = tmp_path / "latin1.txt"
        path.write_bytes("José Pérez".encode("latin-1"))

        with pytest.raises(IngestionError):
            TxtLoader().load(path)

    def test_multiple_newlines_are_preserved(self, tmp_path):
        path = tmp_path / "cv.txt"
        content = "Experience\n\nSkills\n- Python\n- SQL"
        path.write_text(content, encoding="utf-8")

        assert TxtLoader().load(path) == content


class TestMdLoader:
    def test_headings_are_preserved_verbatim(self, tmp_path):
        path = tmp_path / "cv.md"
        path.write_text("## Skills\n- Python\n- SQL", encoding="utf-8")

        result = MdLoader().load(path)

        assert "## Skills" in result

    def test_behaves_like_txt_loader_on_same_input(self, tmp_path):
        path = tmp_path / "cv.md"
        path.write_text("# Title\nBody", encoding="utf-8")

        assert MdLoader().load(path) == TxtLoader().load(path)


@pytest.mark.parametrize(
    "loader_cls", [TxtLoader, MdLoader, DocxLoader, PdfLoader]
)
def test_missing_file_raises_file_not_found(tmp_path, loader_cls):
    missing = tmp_path / "does_not_exist.ext"

    with pytest.raises(FileNotFoundError):
        loader_cls().load(missing)


class TestDocxLoader:
    def test_paragraphs_returned_in_order(self, tmp_path):
        # ING-DOCX-01
        document = docx.Document()
        document.add_paragraph("John Doe")
        document.add_paragraph("Skills: Python, SQL")
        path = tmp_path / "cv.docx"
        document.save(str(path))

        result = DocxLoader().load(path)

        assert result == "John Doe\nSkills: Python, SQL"

    def test_no_paragraphs_returns_empty_string(self, tmp_path):
        # ING-DOCX-02
        document = docx.Document()
        path = tmp_path / "empty.docx"
        document.save(str(path))

        assert DocxLoader().load(path) == ""

    def test_table_content_is_absent_from_result(self, tmp_path):
        # ING-DOCX-03 — documents the known limitation: python-docx's
        # `.paragraphs` does not include text inside tables.
        document = docx.Document()
        document.add_paragraph("John Doe")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Python"
        table.cell(0, 1).text = "SQL"
        path = tmp_path / "with_table.docx"
        document.save(str(path))

        result = DocxLoader().load(path)

        assert "John Doe" in result
        assert "Python" not in result
        assert "SQL" not in result

    def test_corrupted_file_raises_ingestion_error(self, tmp_path):
        # ING-DOCX-04
        path = tmp_path / "corrupted.docx"
        path.write_bytes(b"this is not a real docx file")

        with pytest.raises(IngestionError):
            DocxLoader().load(path)

    def test_blank_paragraphs_are_preserved_as_separators(self, tmp_path):
        # ING-DOCX-05
        document = docx.Document()
        document.add_paragraph("Experience")
        document.add_paragraph("")
        document.add_paragraph("Skills")
        path = tmp_path / "cv.docx"
        document.save(str(path))

        result = DocxLoader().load(path)

        assert result == "Experience\n\nSkills"



class TestLoaderFactory:
    def test_txt_extension_returns_txt_loader(self, tmp_path):
        assert isinstance(
            LoaderFactory.get_loader(tmp_path / "a.txt"), TxtLoader
        )

    def test_unsupported_extension_raises(self, tmp_path):
        with pytest.raises(UnsupportedFileFormatError):
            LoaderFactory.get_loader(tmp_path / "a.rtf")

    def test_uppercase_extension_resolves_same_as_lowercase(self, tmp_path):
        lower = LoaderFactory.get_loader(tmp_path / "a.pdf")
        upper = LoaderFactory.get_loader(tmp_path / "a.PDF")
        assert type(lower) is type(upper)

